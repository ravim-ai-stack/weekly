"""Edits the HMH weekly status report template, which has a different
structure from the other PPTX templates: the cover slide (slide 1) carries
the report date inline in a "Month DD, YYYY" run (no "Week of ..." line on
the content slide at all), and the content slide (slide 2) has 4 relevant
cells: "Completed Platforms & Tasks" (bold sub-header kept in place, bullets
appended below - same shape as the standard template's Completed Tasks),
"Ongoing Work (In Progress)" and "Upcoming Activities ( Next Plan)" (plain
bulleted cells, same shape as the standard template's Upcoming Activities),
and "Risk & Issues" (plain text). A 5th table ("Platforms:" / additional
completed-item narrative) has no identifiable per-week structure and is left
untouched, like every other slide/shape. Every item handed in is kept in
full - none are dropped or truncated to fit available space."""

import re
from datetime import datetime

from pptx import Presentation
from pptx.oxml.ns import qn

from .pptx_builder import _replace_across_runs, set_cell_simple_text, set_list_cell

COVER_SLIDE_INDEX = 0  # slide 1, 0-indexed
CONTENT_SLIDE_INDEX = 1  # slide 2, 0-indexed

COVER_DATE_RE = re.compile(r"[A-Za-z]+\s+\d{1,2},\s+\d{4}")

COMPLETED_HEADER_TEXT = "Completed Platforms & Tasks"


def _normalize_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def _find_table_shape(slide, header_contains: str):
    target = _normalize_ws(header_contains).lower()
    for shape in slide.shapes:
        if shape.has_table:
            try:
                header_text = _normalize_ws(shape.table.cell(0, 0).text)
            except Exception:
                continue
            if target in header_text.lower():
                return shape
    raise RuntimeError(f"Could not find a table containing header '{header_contains}' on the HMH status slide.")


def _set_cover_date(prs, date_text: str):
    slide = prs.slides[COVER_SLIDE_INDEX]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        txBody = shape.text_frame._txBody
        for p in txBody.findall(qn("a:p")):
            if _replace_across_runs(p, COVER_DATE_RE, date_text):
                return
    raise RuntimeError("Could not find the cover date on slide 1.")


def _milestone_line(milestone: dict) -> str:
    area = str(milestone.get("area", "")).strip()
    task = str(milestone.get("task", "")).strip()
    remarks = str(milestone.get("remarks", "")).strip()
    label = " - ".join(p for p in (area, task) if p)
    if label and remarks:
        return f"{label}: {remarks}"
    return remarks or label


def build_hmh_report(template_path: str, output_path: str, week: dict, report: dict):
    """Opens template_path, updates the cover date on slide 1 and the
    Completed Platforms & Tasks / Ongoing Work (In Progress) / Upcoming
    Activities / Risk & Issues cells on slide 2. Every other slide/shape -
    including the Platforms/additional-completed table - is left untouched.
    template_path is never modified."""
    prs = Presentation(template_path)

    cover_date_text = datetime.strptime(week["end"], "%Y-%m-%d").strftime("%B %d, %Y")
    _set_cover_date(prs, cover_date_text)

    slide = prs.slides[CONTENT_SLIDE_INDEX]

    completed_shape = _find_table_shape(slide, "Highlights & Accomplishments")
    set_list_cell(
        completed_shape.table.cell(1, 0),
        report.get("completed_tasks", []),
        header_text=COMPLETED_HEADER_TEXT,
    )

    ongoing_shape = _find_table_shape(slide, "Ongoing Work")
    ongoing_lines = [_milestone_line(m) for m in report.get("milestones", [])]
    set_list_cell(ongoing_shape.table.cell(1, 0), ongoing_lines, header_text=None)

    upcoming_shape = _find_table_shape(slide, "Upcoming Activities")
    set_list_cell(
        upcoming_shape.table.cell(1, 0),
        report.get("upcoming_activities", []),
        header_text=None,
    )

    risk_shape = _find_table_shape(slide, "Risk & Issues")
    risks_text = report.get("risks_issues") or "None"
    set_cell_simple_text(risk_shape.table.cell(1, 0), risks_text)

    prs.save(output_path)
    return output_path
